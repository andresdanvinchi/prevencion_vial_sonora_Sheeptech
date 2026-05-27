from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Presentacion_Prevencion_Vial_Sonora.pptx"
ASSETS = ROOT / "presentacion_assets"


COLORS = {
    "ink": RGBColor(24, 45, 67),
    "muted": RGBColor(87, 101, 116),
    "accent": RGBColor(32, 126, 118),
    "accent2": RGBColor(216, 126, 54),
    "paper": RGBColor(248, 250, 247),
    "line": RGBColor(216, 224, 230),
    "white": RGBColor(255, 255, 255),
}


def add_bg(slide, prs, section=""):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS["paper"]
    bg.line.fill.background()
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["accent"]
    bar.line.fill.background()

    if section:
        tx = slide.shapes.add_textbox(Inches(0.42), Inches(0.18), Inches(4.2), Inches(0.25))
        p = tx.text_frame.paragraphs[0]
        p.text = section.upper()
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = COLORS["accent"]


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(12), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = COLORS["ink"]
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(1.12), Inches(11.6), Inches(0.35))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(12)
        sp.font.color.rgb = COLORS["muted"]


def add_footer(slide, n):
    tx = slide.shapes.add_textbox(Inches(11.75), Inches(7.04), Inches(0.9), Inches(0.2))
    p = tx.text_frame.paragraphs[0]
    p.text = f"{n:02d}"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = COLORS["muted"]
    p.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, bullets, x=0.72, y=1.55, w=4.4, h=4.7, font_size=17):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(8)


def add_image(slide, path, x, y, w, h):
    path = Path(path)
    if not path.exists():
        return
    with Image.open(path) as img:
        iw, ih = img.size
    box_w, box_h = Inches(w), Inches(h)
    ratio = min(box_w / iw, box_h / ih)
    pic_w, pic_h = int(iw * ratio), int(ih * ratio)
    left = Inches(x) + int((box_w - pic_w) / 2)
    top = Inches(y) + int((box_h - pic_h) / 2)
    slide.shapes.add_picture(str(path), left, top, width=pic_w, height=pic_h)


def add_callout(slide, text, x, y, w, h, color="accent"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[color]
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.CENTER


def add_metric_row(slide, items, y=5.8):
    x = 0.72
    for value, label, color in items:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2.45), Inches(0.72))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS[color]
        shape.line.fill.background()
        tf = shape.text_frame
        tf.clear()
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = COLORS["white"]
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = COLORS["white"]
        p2.alignment = PP_ALIGN.CENTER
        x += 2.65


def add_simple_diagram(slide, labels, x=1.1, y=2.35):
    for i, label in enumerate(labels):
        sx = x + i * 2.65
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(sx), Inches(y), Inches(1.85), Inches(0.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["white"]
        shape.line.color.rgb = COLORS["line"]
        tf = shape.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLORS["ink"]
        p.alignment = PP_ALIGN.CENTER
        if i < len(labels) - 1:
            line = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(sx + 1.92), Inches(y + 0.22), Inches(0.55), Inches(0.24))
            line.fill.solid()
            line.fill.fore_color.rgb = COLORS["accent2"]
            line.line.fill.background()


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides = [
        ("Portada", "Análisis de accidentes de tránsito en Sonora para identificar patrones de riesgo vial", ["Minería de Datos", "ATUS INEGI 2015-2024", "Enfoque preventivo"], None),
        ("Problema", "Los accidentes viales tienen patrones que pueden analizarse", ["Generan daños materiales, lesiones y muertes.", "El riesgo no se distribuye igual entre regiones y municipios.", "Analizar solo el total estatal puede ocultar focos de gravedad."], None),
        ("Objetivo", "¿Qué patrones explican el riesgo vial en Sonora?", ["Identificar patrones de frecuencia, gravedad, causa, horario y municipio.", "Usar minería de datos para orientar prevención vial."], "diagram"),
        ("Datos", "Se utilizaron datos ATUS de INEGI 2015-2024", ["Fuente oficial: Accidentes de Tránsito Terrestre en Zonas Urbanas y Suburbanas.", "Filtro: Sonora, clave de entidad 26."], ASSETS / "generated/dataset_resumen.png"),
        ("Preparación", "Antes de analizar, los datos se limpiaron y prepararon", ["Unión de archivos anuales.", "Filtro estatal y eliminación de 'Certificado cero'.", "Normalización de claves, fechas y variables numéricas."], ASSETS / "generated/limpieza_flujo.png"),
        ("Variables", "Se crearon variables para medir riesgo y gravedad", ["La gravedad se definió por heridos o fallecidos.", "Las regiones permitieron comparar zonas con movilidad distinta."], ASSETS / "generated/variables_creadas.png"),
        ("EDA", "El análisis exploratorio mostró volumen, gravedad y causas dominantes", ["Dataset final: 183,911 accidentes.", "16.99% presentan heridos o fallecidos."], ASSETS / "generated/eda_general_con_gravedad.png"),
        ("Regiones", "Hermosillo concentra más accidentes, pero la gravedad cambia por región", ["Hermosillo: mayor volumen.", "Sur y Valle y Frontera también concentran muchos casos.", "Sierra: mayor porcentaje de accidentes graves."], ASSETS / "generated/region_municipio_riesgo.png"),
        ("Tipo y Causa", "La colisión con vehículo automotor domina el registro", ["Tipo más común: colisión con vehículo automotor.", "Causa probable más frecuente: Conductor.", "La prevención debe mirar conducta vial y puntos de conflicto."], ASSETS / "generated/tipo_causa_ranking.png"),
        ("Tiempo", "La tarde concentra el riesgo, con viernes como franja crítica", ["La tarde es la franja con mayor concentración de accidentes.", "Viernes por la tarde aparece como horario crítico: 10,802 accidentes.", "La serie mensual muestra aumento posterior a 2020."], ASSETS / "notebook_outputs/main_cell_23_img_2.png"),
        ("Estadística", "Las variables categóricas sí muestran asociación estadística", ["Chi-cuadrada detectó dependencia entre variables.", "Cramér’s V ayudó a medir la fuerza de asociación."], ASSETS / "generated/chi_cramers_table.png"),
        ("Apriori", "Los atropellamientos se asocian fuertemente con gravedad alta", ["Las reglas de asociación señalaron combinaciones de alto riesgo.", "El hallazgo tiene utilidad directa para proteger peatones."], ASSETS / "generated/apriori_reglas.png"),
        ("Modelos", "Se entrenaron modelos para clasificar accidentes graves", ["Objetivo: GRAVE_BIN.", "Modelos: Árbol de Decisión, Regresión Logística y Random Forest.", "Se evitó fuga de información."], ASSETS / "generated/modelos_pipeline_gravedad.png"),
        ("Comparación", "Random Forest logró el mejor equilibrio para gravedad", ["F1 grave: 0.6298.", "ROC-AUC: 0.8911.", "Regresión Logística tuvo mayor recall grave."], ASSETS / "generated/modelos_comparacion.png"),
        ("Importancia", "El tipo de accidente fue una señal clave para el modelo", ["Variables importantes: tipo de accidente, atropellamiento, motocicleta, vehículos, edad y región.", "La importancia orienta acciones preventivas."], ASSETS / "notebook_outputs/e3_cell_20_img_1.png"),
        ("Clustering", "El clustering agrupó municipios con comportamientos similares", ["Agrupación por volumen, gravedad, heridos, fallecidos y vehículos.", "Permite diseñar estrategias diferenciadas por municipio."], ASSETS / "notebook_outputs/e3_cell_22_img_1.png"),
        ("Prevención", "Los hallazgos se traducen en acciones preventivas", ["Alto volumen: Hermosillo, Sur y Valle, Frontera.", "Alta gravedad: Sierra, Centro, Sur y Valle.", "Riesgos clave: atropellamientos, motocicletas y viernes por la tarde."], ASSETS / "generated/matriz_preventiva.png"),
        ("Conclusiones", "El riesgo vial en Sonora combina volumen y gravedad", ["Volumen y severidad no siempre coinciden.", "Los modelos confirman señales útiles para prevención.", "La minería de datos ayuda a priorizar zonas y condiciones de riesgo."], None),
        ("Recomendaciones", "Recomendaciones para prevención vial basada en datos", ["Priorizar regiones de alto volumen.", "Atender zonas con alta gravedad proporcional.", "Proteger peatones y motociclistas.", "Reforzar acciones viernes por la tarde.", "Complementar ATUS con clima, infraestructura y carreteras."], None),
        ("Referencias", "Referencias", ["INEGI. Accidentes de Tránsito Terrestre en Zonas Urbanas y Suburbanas (ATUS), 1997-2024.", "Pedregosa et al. (2011). Scikit-learn.", "McKinney (2010). pandas.", "Hunter (2007). matplotlib.", "Waskom (2021). seaborn.", "Notebooks E2 y E3 del proyecto."], None),
    ]

    for idx, (section, title, bullets, image) in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        add_bg(slide, prs, section)
        if idx == 1:
            add_title(slide, title)
            add_bullets(slide, bullets, x=0.78, y=2.1, w=6.8, h=1.7, font_size=20)
            add_metric_row(slide, [("2015-2024", "periodo", "accent"), ("183,911", "accidentes reales", "ink"), ("Sonora", "entidad 26", "accent2")], y=5.7)
        elif image:
            add_title(slide, title)
            add_bullets(slide, bullets, x=0.68, y=1.62, w=3.9, h=4.4, font_size=15)
            if idx == 10:
                add_image(slide, ASSETS / "notebook_outputs/main_cell_51_img_1.png", x=4.75, y=1.27, w=7.85, h=2.08)
                add_image(slide, image, x=4.75, y=3.48, w=7.85, h=2.95)
            else:
                add_image(slide, image, x=4.75, y=1.35, w=7.85, h=5.35)
            if idx == 10:
                add_callout(slide, "Viernes tarde: 10,802 accidentes", 0.9, 5.95, 3.35, 0.58, "accent2")
            if idx == 12:
                add_callout(slide, "Atropellamiento -> gravedad alta", 0.9, 5.65, 3.3, 0.58, "accent2")
        else:
            add_title(slide, title)
            if image == "diagram":
                add_bullets(slide, bullets, x=0.85, y=1.7, w=10.8, h=1.2, font_size=18)
                add_simple_diagram(slide, ["Datos ATUS", "Limpieza", "Análisis", "Modelos", "Prevención"], x=0.95, y=3.45)
            else:
                add_bullets(slide, bullets, x=1.05, y=1.65, w=10.6, h=4.7, font_size=20 if idx in [18, 19] else 16)
        add_footer(slide, idx)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_deck()
