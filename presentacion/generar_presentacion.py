from __future__ import annotations

from pathlib import Path
import re
import textwrap
import zipfile

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns

from scipy.stats import chi2_contingency
from scipy.cluster.hierarchy import linkage, dendrogram
from pandas.api.types import is_numeric_dtype
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, roc_auc_score
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier

from mlxtend.frequent_patterns import apriori, association_rules

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentacion"
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "Analisis_Accidentes_Sonora_Riesgo_Vial.pptx"
GUIDE_PATH = OUT_DIR / "guion_presentacion.md"

DATA_DIR = ROOT / "datos" / "originales" / "conjunto_de_datos"
CAT_DIR = ROOT / "datos" / "originales" / "catalogos"

YEARS = range(2015, 2025)
SONORA_ID = "26"

INK = RGBColor(26, 43, 60)
MUTED = RGBColor(94, 107, 120)
BLUE = RGBColor(31, 92, 140)
ORANGE = RGBColor(224, 112, 52)
RED = RGBColor(180, 55, 55)
GREEN = RGBColor(64, 126, 89)
PAPER = RGBColor(248, 249, 250)
LIGHT_BLUE = RGBColor(226, 237, 246)
LIGHT_ORANGE = RGBColor(253, 237, 226)
WHITE = RGBColor(255, 255, 255)


def clean_id(series: pd.Series, width: int) -> pd.Series:
    return series.astype(str).str.strip().str.replace(r"\.0$", "", regex=True).str.zfill(width)


def load_data() -> pd.DataFrame:
    frames = []
    for year in YEARS:
        df = pd.read_csv(
            DATA_DIR / f"atus_anual_{year}.csv",
            dtype=str,
            low_memory=False,
            index_col=False,
            skipinitialspace=True,
        )
        df.columns = df.columns.str.upper().str.strip()
        df["ID_ENTIDAD"] = clean_id(df["ID_ENTIDAD"], 2)
        df["ID_MUNICIPIO"] = clean_id(df["ID_MUNICIPIO"], 3)
        df["ANIO"] = year
        frames.append(df.loc[df["ID_ENTIDAD"].eq(SONORA_ID)].copy())

    atus = pd.concat(frames, ignore_index=True)
    municipios = pd.read_csv(CAT_DIR / "tc_municipio.csv", dtype=str, low_memory=False, index_col=False, skipinitialspace=True)
    municipios.columns = municipios.columns.str.upper().str.strip()
    municipios["ID_ENTIDAD"] = clean_id(municipios["ID_ENTIDAD"], 2)
    municipios["ID_MUNICIPIO"] = clean_id(municipios["ID_MUNICIPIO"], 3)
    atus = atus.merge(
        municipios[["ID_ENTIDAD", "ID_MUNICIPIO", "NOM_MUNICIPIO"]],
        on=["ID_ENTIDAD", "ID_MUNICIPIO"],
        how="left",
    )

    for col in atus.select_dtypes(include=["object", "string"]).columns:
        atus[col] = atus[col].astype(str).str.strip()
        atus[col] = atus[col].replace({"nan": np.nan, "None": np.nan})

    atus = atus.loc[~atus["TIPACCID"].eq("Certificado cero")].copy()

    dead_cols = ["CONDMUERTO", "PASAMUERTO", "PEATMUERTO", "CICLMUERTO", "OTROMUERTO", "NEMUERTO"]
    injury_cols = ["CONDHERIDO", "PASAHERIDO", "PEATHERIDO", "CICLHERIDO", "OTROHERIDO", "NEHERIDO"]
    vehicle_cols = [
        "AUTOMOVIL", "CAMPASAJ", "MICROBUS", "PASCAMION", "OMNIBUS", "TRANVIA",
        "CAMIONETA", "CAMION", "TRACTOR", "FERROCARRI", "MOTOCICLET", "BICICLETA", "OTROVEHIC",
    ]
    for col in dead_cols + injury_cols + vehicle_cols + ["ID_HORA", "ID_MINUTO", "ID_DIA", "ID_EDAD", "ANIO"]:
        atus[col] = pd.to_numeric(atus[col], errors="coerce")

    atus["TOTAL_MUERTOS"] = atus[dead_cols].fillna(0).sum(axis=1)
    atus["TOTAL_HERIDOS"] = atus[injury_cols].fillna(0).sum(axis=1)
    atus["TOTAL_VICTIMAS"] = atus["TOTAL_MUERTOS"] + atus["TOTAL_HERIDOS"]
    atus["TOTAL_VEHICULOS"] = atus[vehicle_cols].fillna(0).sum(axis=1)
    atus["GRAVE_BIN"] = ((atus["TOTAL_VICTIMAS"] > 0) | atus["CLASACC"].isin(["Fatal", "No fatal"])).astype(int)
    atus["NIVEL_GRAVEDAD"] = np.select(
        [atus["TOTAL_MUERTOS"] > 0, atus["TOTAL_HERIDOS"] > 0],
        ["Fatal", "Con heridos"],
        default="Solo daños",
    )
    atus["INVOLUCRA_MOTO"] = atus["MOTOCICLET"].fillna(0).gt(0).astype(int)
    atus["INVOLUCRA_BICI"] = atus["BICICLETA"].fillna(0).gt(0).astype(int)

    region_map = {
        "030": "Hermosillo",
        "002": "Frontera", "019": "Frontera", "027": "Frontera", "035": "Frontera",
        "039": "Frontera", "041": "Frontera", "043": "Frontera", "055": "Frontera",
        "058": "Frontera", "059": "Frontera", "060": "Frontera", "067": "Frontera",
        "004": "Costa y Desierto", "007": "Costa y Desierto", "017": "Costa y Desierto",
        "046": "Costa y Desierto", "047": "Costa y Desierto", "048": "Costa y Desierto",
        "064": "Costa y Desierto", "065": "Costa y Desierto", "070": "Costa y Desierto",
        "001": "Centro", "006": "Centro", "013": "Centro", "014": "Centro", "016": "Centro",
        "020": "Centro", "021": "Centro", "022": "Centro", "023": "Centro", "024": "Centro",
        "028": "Centro", "034": "Centro", "037": "Centro", "038": "Centro", "045": "Centro",
        "050": "Centro", "053": "Centro", "054": "Centro", "056": "Centro", "057": "Centro",
        "063": "Centro", "066": "Centro", "068": "Centro",
        "003": "Sur y Valle", "012": "Sur y Valle", "018": "Sur y Valle", "025": "Sur y Valle",
        "026": "Sur y Valle", "029": "Sur y Valle", "033": "Sur y Valle", "042": "Sur y Valle",
        "049": "Sur y Valle", "051": "Sur y Valle", "071": "Sur y Valle", "072": "Sur y Valle",
        "005": "Sierra", "008": "Sierra", "009": "Sierra", "010": "Sierra", "011": "Sierra",
        "015": "Sierra", "031": "Sierra", "032": "Sierra", "040": "Sierra", "044": "Sierra",
        "052": "Sierra", "061": "Sierra", "062": "Sierra", "069": "Sierra",
    }
    atus["REGION"] = atus["ID_MUNICIPIO"].map(region_map).fillna("Otros municipios")
    atus["MES_NUM"] = pd.to_numeric(atus["MES"], errors="coerce")
    atus["HORA_VALIDA"] = atus["ID_HORA"].where(atus["ID_HORA"].between(0, 23))
    atus["RANGO_HORA"] = pd.cut(
        atus["HORA_VALIDA"],
        bins=[-0.1, 5, 11, 17, 23],
        labels=["Madrugada", "Mañana", "Tarde", "Noche"],
    )
    atus["EDAD_CONDUCTOR"] = atus["ID_EDAD"].where(atus["ID_EDAD"].between(12, 98))
    atus["DIASEMANA"] = (
        atus["DIASEMANA"]
        .replace({
            "Sabado": "Sábado",
            "Miercoles": "Miércoles",
        })
    )
    return atus


def fig_path(name: str) -> Path:
    return ASSET_DIR / f"{name}.png"


def savefig(name: str) -> Path:
    path = fig_path(name)
    plt.tight_layout()
    plt.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def wrap_label(value: str, width: int = 18) -> str:
    return "\n".join(textwrap.wrap(str(value), width=width))


def create_charts(atus: pd.DataFrame) -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    sns.set_theme(style="whitegrid")
    paths: dict[str, Path] = {}

    annual = atus["ANIO"].value_counts().sort_index().reindex(list(YEARS), fill_value=0)
    plt.figure(figsize=(8.5, 4.3))
    ax = plt.gca()
    ax.plot(
        annual.index,
        annual.values,
        color="#1f5c8c",
        marker="o",
        linewidth=3,
        markersize=7,
    )
    ax.fill_between(annual.index, annual.values, color="#1f5c8c", alpha=0.12)
    ax.set_title("Histórico de accidentes de tránsito en Sonora, 2015-2024", fontsize=14, weight="bold")
    ax.set_xlabel("Año")
    ax.set_ylabel("Accidentes")
    ax.set_xticks(list(YEARS))
    ax.set_ylim(0, max(annual.values) * 1.18)
    for year, value in annual.items():
        if year in [min(YEARS), max(YEARS)] or value == annual.max() or year == 2020:
            ax.annotate(
                f"{int(value):,}",
                xy=(year, value),
                xytext=(0, 9),
                textcoords="offset points",
                ha="center",
                fontsize=8,
                color="#1a2b3c",
            )
    ax.grid(axis="y", alpha=0.35)
    paths["annual"] = savefig("accidentes_por_anio")

    region = atus.groupby("REGION").agg(accidentes=("GRAVE_BIN", "size"), pct_graves=("GRAVE_BIN", "mean"))
    region = region.sort_values("accidentes", ascending=False)
    fig, ax1 = plt.subplots(figsize=(8.5, 4.3))
    x = np.arange(len(region))
    ax1.bar(x, region["accidentes"], color="#e07034")
    ax1.set_xticks(x)
    ax1.set_xticklabels([wrap_label(i, 14) for i in region.index], rotation=0, fontsize=8)
    ax1.set_ylabel("Accidentes")
    ax1.set_title("Accidentes por región y porcentaje de gravedad", fontsize=14, weight="bold")
    ax2 = ax1.twinx()
    ax2.plot(x, region["pct_graves"] * 100, color="#b43737", marker="o", linewidth=2.5)
    ax2.set_ylabel("% graves")
    paths["region"] = savefig("accidentes_region_gravedad")

    top_mun = atus["NOM_MUNICIPIO"].value_counts().head(10).sort_values()
    plt.figure(figsize=(8, 4.5))
    top_mun.plot(kind="barh", color="#407e59")
    plt.title("Municipios con mayor volumen de accidentes", fontsize=14, weight="bold")
    plt.xlabel("Accidentes")
    plt.ylabel("")
    paths["municipios"] = savefig("top_municipios")

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2))
    types = atus["TIPACCID"].value_counts().head(7).sort_values()
    causes = atus["CAUSAACCI"].value_counts().head(6).sort_values()
    types.plot(kind="barh", ax=axes[0], color="#1f5c8c")
    axes[0].set_title("Tipo de accidente")
    axes[0].set_xlabel("Accidentes")
    axes[0].set_ylabel("")
    causes.plot(kind="barh", ax=axes[1], color="#e07034")
    axes[1].set_title("Causa probable")
    axes[1].set_xlabel("Accidentes")
    axes[1].set_ylabel("")
    paths["tipo_causa"] = savefig("tipo_y_causa")

    temporal = pd.crosstab(atus["DIASEMANA"], atus["RANGO_HORA"], values=atus["GRAVE_BIN"], aggfunc="mean") * 100
    day_order = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    temporal = temporal.reindex([d for d in day_order if d in temporal.index])
    plt.figure(figsize=(8, 4.6))
    sns.heatmap(temporal, annot=True, fmt=".1f", cmap="Reds", cbar_kws={"label": "% graves"})
    plt.title("Gravedad por día y rango horario", fontsize=14, weight="bold")
    plt.xlabel("Rango horario")
    plt.ylabel("Día de la semana")
    paths["temporal_heat"] = savefig("heatmap_temporal")

    monthly = atus.groupby("MES_NUM").size().rename_axis("MES_NUM").reset_index(name="accidentes")
    plt.figure(figsize=(8.5, 4.2))
    ax = sns.barplot(data=monthly, x="MES_NUM", y="accidentes", color="#1f5c8c")
    ax.bar_label(ax.containers[0], labels=[f"{int(v/1000)}k" for v in monthly["accidentes"]], fontsize=8, padding=2)
    plt.title("Accidentes acumulados por mes, 2015-2024", fontsize=14, weight="bold")
    plt.xlabel("Mes")
    plt.ylabel("Accidentes")
    paths["monthly"] = savefig("series_mensuales")

    return paths


def cramers_v(x: pd.Series, y: pd.Series) -> tuple[float, float]:
    table = pd.crosstab(x, y)
    chi2, p, _, _ = chi2_contingency(table)
    n = table.values.sum()
    r, k = table.shape
    v = np.sqrt((chi2 / n) / max(1, min(k - 1, r - 1)))
    return float(p), float(v)


def build_analysis(atus: pd.DataFrame, paths: dict[str, Path]) -> dict:
    total = len(atus)
    annual = atus["ANIO"].value_counts().sort_index()
    region_summary = (
        atus.groupby("REGION")
        .agg(accidentes=("GRAVE_BIN", "size"), pct_graves=("GRAVE_BIN", "mean"), heridos=("TOTAL_HERIDOS", "sum"), muertos=("TOTAL_MUERTOS", "sum"))
        .assign(pct_graves=lambda df: (df["pct_graves"] * 100).round(2))
        .sort_values("accidentes", ascending=False)
    )
    top_mun = atus["NOM_MUNICIPIO"].value_counts().head(5)
    top_type = atus["TIPACCID"].value_counts().idxmax()
    top_cause = atus["CAUSAACCI"].value_counts().idxmax()
    critical = (
        atus.groupby(["DIASEMANA", "RANGO_HORA"], observed=False)
        .agg(accidentes=("GRAVE_BIN", "size"), pct_graves=("GRAVE_BIN", "mean"))
        .dropna()
        .sort_values(["pct_graves", "accidentes"], ascending=False)
        .head(1)
    )
    chi_vars = ["REGION", "TIPACCID", "CAUSAACCI", "RANGO_HORA", "DIASEMANA", "ALIENTO", "CINTURON", "SEXO"]
    chi_rows = []
    for var in chi_vars:
        tmp = atus[[var, "GRAVE_BIN"]].dropna()
        p, v = cramers_v(tmp[var], tmp["GRAVE_BIN"])
        chi_rows.append({"variable": var, "p_value": p, "cramers_v": v})
    chi_df = pd.DataFrame(chi_rows).sort_values("cramers_v", ascending=False)

    plt.figure(figsize=(7.5, 4.2))
    plot_chi = chi_df.sort_values("cramers_v").copy()
    chi_label_map = {
        "TIPACCID": "Tipo de accidente",
        "CAUSAACCI": "Causa probable",
        "CINTURON": "Cinturón",
        "DIASEMANA": "Día",
        "REGION": "Región",
        "ALIENTO": "Aliento",
        "RANGO_HORA": "Rango horario",
        "SEXO": "Sexo",
    }
    plt.barh(plot_chi["variable"].map(chi_label_map).fillna(plot_chi["variable"]), plot_chi["cramers_v"], color="#1f5c8c")
    plt.title("Asociación con gravedad: Cramér's V", fontsize=14, weight="bold")
    plt.xlabel("Cramér's V")
    paths["cramers"] = savefig("cramers_v")

    rule_df = atus[["REGION", "TIPACCID", "CAUSAACCI", "RANGO_HORA", "DIASEMANA", "GRAVE_BIN"]].dropna().copy()
    rule_df["GRAVEDAD"] = np.where(rule_df["GRAVE_BIN"].eq(1), "Grave", "No grave")
    rule_df = rule_df.drop(columns=["GRAVE_BIN"])
    for col in rule_df.columns:
        rule_df[col] = col + "=" + rule_df[col].astype(str)
    sample_rules = rule_df.sample(n=min(50000, len(rule_df)), random_state=42)
    one_hot = pd.get_dummies(sample_rules)
    freq = apriori(one_hot, min_support=0.015, use_colnames=True)
    rules = association_rules(freq, metric="confidence", min_threshold=0.25)
    rules = rules[rules["consequents"].astype(str).str.contains("GRAVEDAD=Grave", regex=False)].copy()
    if not rules.empty:
        rules["antecedent_text"] = rules["antecedents"].apply(lambda s: ", ".join(sorted(str(i).split("_", 1)[-1] for i in s)))
        rules = rules.sort_values(["lift", "confidence", "support"], ascending=False).head(5)
    else:
        rules = pd.DataFrame(columns=["antecedent_text", "support", "confidence", "lift"])

    model_df = atus.dropna(subset=["GRAVE_BIN"]).copy()
    model_df = model_df.sample(n=min(60000, len(model_df)), random_state=42)
    features = ["REGION", "TIPACCID", "CAUSAACCI", "RANGO_HORA", "DIASEMANA", "SEXO", "ALIENTO", "CINTURON", "CAPAROD", "TOTAL_VEHICULOS", "INVOLUCRA_MOTO", "INVOLUCRA_BICI", "HORA_VALIDA", "EDAD_CONDUCTOR"]
    model_df = model_df[features + ["GRAVE_BIN"]].dropna()
    X = model_df[features]
    y = model_df["GRAVE_BIN"]
    cat_cols = [c for c in features if not is_numeric_dtype(X[c])]
    num_cols = [c for c in features if c not in cat_cols]
    preprocess = ColumnTransformer(
        transformers=[
            ("cat", OneHotEncoder(handle_unknown="ignore", min_frequency=30), cat_cols),
            ("num", Pipeline([("scaler", StandardScaler())]), num_cols),
        ]
    )
    X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.25, random_state=42, stratify=y)
    models = {
        "Arbol de decision": DecisionTreeClassifier(max_depth=8, random_state=42, class_weight="balanced"),
        "Regresion logistica": LogisticRegression(max_iter=800, class_weight="balanced"),
        "Random Forest": RandomForestClassifier(n_estimators=120, max_depth=12, min_samples_leaf=8, random_state=42, n_jobs=-1, class_weight="balanced_subsample"),
    }
    rows = []
    fitted = {}
    for name, estimator in models.items():
        pipe = Pipeline([("preprocess", preprocess), ("model", estimator)])
        pipe.fit(X_train, y_train)
        pred = pipe.predict(X_test)
        proba = pipe.predict_proba(X_test)[:, 1] if hasattr(pipe.named_steps["model"], "predict_proba") else pred
        rows.append({
            "Modelo": name,
            "Accuracy": accuracy_score(y_test, pred),
            "Precision": precision_score(y_test, pred, zero_division=0),
            "Recall": recall_score(y_test, pred, zero_division=0),
            "F1": f1_score(y_test, pred, zero_division=0),
            "ROC-AUC": roc_auc_score(y_test, proba),
        })
        fitted[name] = pipe
    model_metrics = pd.DataFrame(rows).sort_values("F1", ascending=False)

    plt.figure(figsize=(8, 4.4))
    metric_plot = model_metrics.set_index("Modelo")[["Precision", "Recall", "F1", "ROC-AUC"]].rename(columns={"Precision": "Precisión"})
    metric_plot.plot(kind="bar", ax=plt.gca(), color=["#1f5c8c", "#e07034", "#407e59", "#b43737"])
    plt.title("Comparación de modelos de clasificación", fontsize=14, weight="bold")
    plt.ylabel("Valor")
    plt.xlabel("")
    plt.ylim(0, 1)
    plt.xticks(rotation=15, ha="right")
    plt.legend(loc="upper center", bbox_to_anchor=(0.5, -0.18), ncol=4, fontsize=8, frameon=False)
    paths["models"] = savefig("comparacion_modelos")

    rf = fitted["Random Forest"]
    feature_names = rf.named_steps["preprocess"].get_feature_names_out()
    importances = rf.named_steps["model"].feature_importances_
    importance = pd.DataFrame({"feature": feature_names, "importance": importances})
    importance["feature_clean"] = (
        importance["feature"]
        .str.replace(r"^(cat|num)__", "", regex=True)
        .str.replace("TIPACCID_", "Tipo: ", regex=False)
        .str.replace("CAUSAACCI_", "Causa: ", regex=False)
        .str.replace("REGION_", "Región: ", regex=False)
        .str.replace("RANGO_HORA_", "Horario: ", regex=False)
        .str.replace("DIASEMANA_", "Día: ", regex=False)
        .str.replace("CINTURON_", "Cinturón: ", regex=False)
        .str.replace("TOTAL_VEHICULOS", "Total vehículos", regex=False)
        .str.replace("INVOLUCRA_MOTO", "Involucra moto", regex=False)
        .str.replace("INVOLUCRA_BICI", "Involucra bici", regex=False)
        .str.replace("HORA_VALIDA", "Hora válida", regex=False)
        .str.replace("EDAD_CONDUCTOR", "Edad conductor", regex=False)
        .str.replace("_", " ")
    )
    importance = importance.sort_values("importance", ascending=False).head(12)
    plt.figure(figsize=(8, 4.6))
    plt.barh(importance["feature_clean"][::-1], importance["importance"][::-1], color="#1f5c8c")
    plt.title("Importancia de variables - Random Forest", fontsize=14, weight="bold")
    plt.xlabel("Importancia")
    paths["importance"] = savefig("importancia_variables")

    municipal = (
        atus.groupby(["NOM_MUNICIPIO", "REGION"])
        .agg(accidentes=("GRAVE_BIN", "size"), pct_graves=("GRAVE_BIN", "mean"), motos=("INVOLUCRA_MOTO", "mean"), peatones=("TIPACCID", lambda s: s.astype(str).str.contains("peat", case=False, na=False).mean()))
        .reset_index()
    )
    municipal = municipal[municipal["accidentes"] >= 50].copy()
    municipal = municipal.sort_values("accidentes", ascending=False).head(30).copy()
    cluster_features = municipal[["accidentes", "pct_graves", "motos", "peatones"]].copy()
    cluster_features = (cluster_features - cluster_features.mean()) / cluster_features.std(ddof=0)
    Z = linkage(cluster_features.fillna(0), method="ward")
    plt.figure(figsize=(9, 4.7))
    dendrogram(Z, labels=municipal["NOM_MUNICIPIO"].values, leaf_rotation=90, leaf_font_size=6, color_threshold=None)
    plt.title("Clustering jerárquico de municipios", fontsize=14, weight="bold")
    plt.ylabel("Distancia")
    paths["cluster"] = savefig("clustering_municipios")

    return {
        "total": total,
        "annual": annual,
        "region": region_summary,
        "top_mun": top_mun,
        "top_type": top_type,
        "top_cause": top_cause,
        "critical": critical,
        "chi": chi_df,
        "rules": rules,
        "model_metrics": model_metrics,
        "importance": importance,
        "municipal": municipal,
    }


def add_textbox(slide, x, y, w, h, text, font_size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, fill=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = fill
    return shape


def add_title(slide, kicker: str, title: str):
    add_textbox(slide, 0.55, 0.28, 2.0, 0.28, kicker.upper(), 9, ORANGE, True)
    add_textbox(slide, 0.55, 0.55, 8.8, 0.55, title, 24, INK, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.14), Inches(1.0), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.color.rgb = ORANGE


def add_footer(slide, n: int):
    add_textbox(slide, 0.55, 7.05, 8.5, 0.18, "Fuente: E1, ATUS / INEGI 2015-2024, notebooks E2 y E3", 7, MUTED)
    add_textbox(slide, 12.15, 7.02, 0.55, 0.2, f"{n:02d}", 8, MUTED, True, PP_ALIGN.RIGHT)


def add_bullets(slide, x, y, w, h, bullets, font_size=15, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def add_card(slide, x, y, w, h, value, label, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(220, 225, 230)
    add_textbox(slide, x + 0.18, y + 0.18, w - 0.36, 0.36, value, 22, color, True, PP_ALIGN.CENTER)
    add_textbox(slide, x + 0.18, y + 0.62, w - 0.36, 0.42, label, 9, MUTED, False, PP_ALIGN.CENTER)


def add_image(slide, path: Path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def make_table(slide, df: pd.DataFrame, x, y, w, h, font_size=8):
    rows, cols = df.shape[0] + 1, df.shape[1]
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
        table.cell(0, j).fill.solid()
        table.cell(0, j).fill.fore_color.rgb = LIGHT_BLUE
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, value in enumerate(row):
            table.cell(i, j).text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.color.rgb = INK
    return table


def clean_metric(v: float) -> str:
    return f"{v:.2f}" if isinstance(v, float) else str(v)


def build_deck(atus: pd.DataFrame, analysis: dict, paths: dict[str, Path]):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides_meta = []

    def new_slide(kicker, title):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = PAPER
        add_title(slide, kicker, title)
        add_footer(slide, len(prs.slides))
        slides_meta.append({"title": title, "notes": ""})
        return slide

    # 1
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    add_textbox(slide, 0.75, 0.7, 11.5, 1.6, "Análisis de accidentes de tránsito en Sonora para identificar patrones de riesgo vial", 32, WHITE, True)
    add_textbox(slide, 0.78, 2.28, 6.0, 0.45, "Proyecto de Minería de Datos | ATUS / INEGI 2015-2024", 16, RGBColor(223, 232, 240))
    add_card(slide, 0.82, 3.35, 2.4, 1.1, f"{analysis['total']:,}", "accidentes analizados", ORANGE)
    add_card(slide, 3.45, 3.35, 2.0, 1.1, "2015-2024", "periodo de estudio", BLUE)
    add_card(slide, 5.68, 3.35, 2.1, 1.1, "Prevención", "enfoque final", GREEN)
    add_image(slide, paths["annual"], 8.15, 2.2, 4.55, 2.7)
    add_textbox(slide, 0.8, 6.65, 10.5, 0.25, "Integrantes: [nombres]   |   Materia: [materia]   |   Fecha: [fecha]", 10, RGBColor(210, 218, 225))
    slides_meta.append({"title": "Portada", "notes": "Presentar el tema, la fuente de datos y el enfoque preventivo del proyecto."})

    # 2
    slide = new_slide("Problema", "Los accidentes viales requieren análisis con enfoque preventivo")
    add_bullets(slide, 0.75, 1.45, 4.25, 3.4, [
        "Los accidentes generan daños materiales, lesiones y fallecimientos.",
        "El riesgo no se distribuye igual por municipio, horario o tipo de accidente.",
        "La minería de datos ayuda a convertir registros históricos en evidencia para decidir mejor.",
    ], 16)
    add_image(slide, paths["annual"], 5.35, 1.45, 7.25, 4.25)

    # 3
    slide = new_slide("Objetivo", "La pregunta guía conecta datos históricos con prevención vial")
    add_textbox(slide, 0.85, 1.45, 5.65, 1.05, "Pregunta de investigación", 18, ORANGE, True)
    add_textbox(slide, 0.85, 1.95, 5.7, 1.2, "¿Qué patrones de riesgo vial pueden identificarse en los accidentes de tránsito de Sonora entre 2015 y 2024?", 20, INK, True)
    add_textbox(slide, 7.0, 1.45, 5.2, 1.05, "Objetivo general", 18, ORANGE, True)
    add_bullets(slide, 7.0, 1.95, 5.15, 2.8, [
        "Analizar datos ATUS de Sonora.",
        "Detectar patrones por región, gravedad, tipo de accidente y tiempo.",
        "Evaluar técnicas de minería de datos con utilidad preventiva.",
    ], 16)

    # 4
    slide = new_slide("Dataset", "ATUS / INEGI permite estudiar accidentes con cobertura anual")
    add_card(slide, 0.8, 1.45, 2.3, 1.05, "10 años", "2015 a 2024", BLUE)
    add_card(slide, 3.35, 1.45, 2.7, 1.05, f"{analysis['total']:,}", "registros finales", ORANGE)
    add_card(slide, 6.3, 1.45, 2.25, 1.05, "Sonora", "entidad 26", GREEN)
    add_card(slide, 8.8, 1.45, 2.5, 1.05, "INEGI", "fuente oficial", RED)
    add_bullets(slide, 0.95, 3.05, 5.4, 2.3, [
        "Se integraron archivos anuales ATUS.",
        "Se filtraron registros de Sonora.",
        "Se añadió el catálogo municipal para ubicar municipios.",
        "El análisis se documenta en E2 y E3.",
    ], 15)
    add_image(slide, paths["annual"], 6.8, 3.0, 5.55, 3.0)

    # 5
    slide = new_slide("Preparación", "La limpieza deja una base confiable para el análisis")
    add_bullets(slide, 0.78, 1.45, 5.2, 4.1, [
        "Unificación de CSV anuales 2015-2024.",
        "Filtro de entidad Sonora con clave 26.",
        "Eliminación de registros 'Certificado cero'.",
        "Conversión de columnas numéricas.",
        "Creación de variables temporales, regionales y de gravedad.",
    ], 16)
    prep_df = pd.DataFrame({
        "Paso": ["Carga", "Filtro", "Depuración", "Variables"],
        "Resultado": ["10 CSV anuales", "Entidad 26", "Accidentes reales", "Base para EDA/E3"],
    })
    make_table(slide, prep_df, 6.4, 1.65, 5.65, 1.8, 11)
    add_textbox(slide, 6.45, 4.0, 5.4, 0.85, "Frase clave: la calidad del análisis depende primero de una base limpia y consistente.", 17, INK, True)

    # 6
    slide = new_slide("Variables", "Región, gravedad y víctimas traducen los datos a riesgo vial")
    var_df = pd.DataFrame({
        "Variable": ["REGION", "TOTAL_VICTIMAS", "GRAVE_BIN", "RANGO_HORA", "INVOLUCRA_MOTO/BICI"],
        "Uso": ["Segmentar Sonora", "Heridos + fallecidos", "Grave si hubo víctimas", "Madrugada, mañana, tarde, noche", "Identificar usuarios vulnerables"],
    })
    make_table(slide, var_df, 0.85, 1.5, 7.0, 2.6, 10)
    add_bullets(slide, 8.25, 1.55, 4.0, 3.0, [
        "La gravedad se definió por presencia de personas heridas o fallecidas.",
        "La región permite comparar patrones territoriales.",
        "Las variables creadas conectan el análisis con decisiones preventivas.",
    ], 15)

    # 7
    slide = new_slide("EDA", "El análisis exploratorio muestra crecimiento reciente y concentración territorial")
    add_image(slide, paths["annual"], 0.75, 1.45, 5.7, 3.35)
    add_image(slide, paths["region"], 6.75, 1.45, 5.8, 3.35)
    add_textbox(slide, 0.95, 5.25, 11.4, 0.6, f"Hallazgo: el pico anual aparece en {analysis['annual'].idxmax()} con {analysis['annual'].max():,} accidentes.", 18, INK, True, PP_ALIGN.CENTER)

    # 8
    slide = new_slide("Territorio", "Hermosillo concentra el mayor volumen; Sierra destaca por gravedad")
    add_image(slide, paths["region"], 0.75, 1.42, 6.4, 3.7)
    add_image(slide, paths["municipios"], 7.35, 1.42, 5.1, 3.7)
    add_bullets(slide, 0.9, 5.45, 11.1, 0.95, [
        f"Hermosillo encabeza el volumen regional con {int(analysis['region'].loc['Hermosillo','accidentes']):,} accidentes.",
        "Sur y Valle y Frontera también concentran una parte importante del problema vial.",
        "La Sierra requiere atención porque su proporción de accidentes graves es alta.",
    ], 12)

    # 9
    slide = new_slide("Tipo y causa", "La colisión vehicular domina y la causa principal se atribuye al conductor")
    add_image(slide, paths["tipo_causa"], 0.75, 1.4, 8.0, 4.55)
    add_bullets(slide, 9.05, 1.55, 3.4, 3.6, [
        f"Tipo más común: {analysis['top_type']}.",
        f"Causa probable más frecuente: {analysis['top_cause']}.",
        "Los atropellamientos se relacionan con mayor gravedad.",
    ], 14)

    # 10
    slide = new_slide("Tiempo", "El horario permite ubicar ventanas críticas de prevención")
    add_image(slide, paths["temporal_heat"], 0.75, 1.35, 5.75, 3.85)
    add_image(slide, paths["monthly"], 6.8, 1.35, 5.75, 3.85)
    friday_afternoon = atus.loc[atus["DIASEMANA"].astype(str).str.contains("Viernes", case=False, na=False) & atus["RANGO_HORA"].astype(str).eq("Tarde")]
    add_textbox(slide, 0.95, 5.55, 11.15, 0.5, f"Horario crítico preventivo: viernes por la tarde ({len(friday_afternoon):,} accidentes en el periodo).", 17, INK, True, PP_ALIGN.CENTER)

    # 11
    slide = new_slide("Asociación", "Chi-cuadrada y Cramér's V miden qué variables se relacionan con gravedad")
    add_image(slide, paths["cramers"], 0.85, 1.45, 6.0, 3.9)
    top_chi = analysis["chi"].head(5).copy()
    top_chi["p_value"] = top_chi["p_value"].map(lambda v: "<0.001" if v < 0.001 else f"{v:.3f}")
    top_chi["cramers_v"] = top_chi["cramers_v"].map(lambda v: f"{v:.3f}")
    make_table(slide, top_chi.rename(columns={"variable": "Variable", "p_value": "p", "cramers_v": "V"}), 7.25, 1.6, 4.9, 2.25, 10)
    add_textbox(slide, 7.25, 4.35, 4.8, 0.9, "Interpretación: valores mayores de V indican asociación más fuerte con la gravedad; no significan causalidad directa.", 14, INK)

    # 12
    slide = new_slide("Apriori", "Las reglas de asociación encuentran combinaciones frecuentes de riesgo")
    rules_show = analysis["rules"][["antecedent_text", "support", "confidence", "lift"]].copy()
    if rules_show.empty:
        rules_show = pd.DataFrame({"antecedente": ["Sin reglas con el umbral usado"], "support": [""], "confidence": [""], "lift": [""]})
    else:
        rules_show = rules_show.rename(columns={"antecedent_text": "Antecedente", "support": "Soporte", "confidence": "Confianza", "lift": "Lift"})
        for col in ["Soporte", "Confianza", "Lift"]:
            rules_show[col] = rules_show[col].map(lambda v: f"{v:.2f}")
        rules_show["Antecedente"] = rules_show["Antecedente"].str.slice(0, 70)
    make_table(slide, rules_show, 0.65, 1.45, 12.0, 2.9, 8)
    add_bullets(slide, 1.0, 4.75, 10.9, 1.0, [
        "Apriori no predice por sí solo; ayuda a descubrir patrones recurrentes.",
        "Las reglas son útiles para traducir combinaciones de condiciones en mensajes preventivos.",
    ], 14)

    # 13
    slide = new_slide("Clasificación", "Los modelos buscan anticipar si un accidente será grave")
    add_bullets(slide, 0.85, 1.45, 5.2, 3.35, [
        "Variable objetivo: GRAVE_BIN.",
        "Modelos: árbol de decisión, regresión logística y Random Forest.",
        "Se evita fuga de información quitando heridos, muertos y clasificación original como predictores.",
        "La evaluación considera precisión, recall, F1 y ROC-AUC.",
    ], 15)
    add_textbox(slide, 6.65, 1.75, 4.8, 1.1, "En prevención, el recall importa porque interesa detectar la mayor cantidad posible de casos graves.", 22, INK, True, PP_ALIGN.CENTER)
    add_card(slide, 7.2, 3.45, 3.6, 1.05, "Random Forest", "modelo fuerte e interpretable por importancia", BLUE)

    # 14
    slide = new_slide("Evaluación", "Random Forest ofrece el mejor equilibrio entre desempeño e interpretación")
    add_image(slide, paths["models"], 0.75, 1.35, 7.0, 4.25)
    metrics = analysis["model_metrics"].copy()
    for c in ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]:
        metrics[c] = metrics[c].map(lambda v: f"{v:.2f}")
    metrics = metrics.rename(columns={"Precision": "Precisión"})
    make_table(slide, metrics, 8.0, 1.55, 4.65, 2.4, 8)
    add_textbox(slide, 8.05, 4.55, 4.45, 0.8, "La comparación no solo busca exactitud: también revisa sensibilidad ante accidentes graves.", 14, INK)

    # 15
    slide = new_slide("Importancia", "Random Forest ayuda a explicar qué señales pesan más en la gravedad")
    add_image(slide, paths["importance"], 0.85, 1.35, 7.0, 4.55)
    add_bullets(slide, 8.25, 1.55, 3.8, 3.4, [
        "La importancia de variables permite interpretar el modelo.",
        "Tipo de accidente, causa, horario y región aportan señales preventivas.",
        "No se usa como prueba causal, sino como apoyo para priorizar análisis.",
    ], 15)

    # 16
    slide = new_slide("Clustering", "El agrupamiento compara municipios con comportamientos similares")
    add_image(slide, paths["cluster"], 0.75, 1.35, 7.45, 4.65)
    add_bullets(slide, 8.5, 1.55, 3.8, 3.4, [
        "El clustering jerárquico agrupa municipios por volumen, gravedad y patrones de usuarios vulnerables.",
        "Sirve para comparar municipios aunque tengan tamaños distintos.",
        "Permite pensar estrategias por perfiles de riesgo, no solo por ubicación.",
    ], 14)

    # 17
    slide = new_slide("Prevención", "Los resultados deben convertirse en decisiones preventivas")
    add_bullets(slide, 0.9, 1.45, 5.45, 4.0, [
        "Hermosillo: priorizar gestión de volumen y puntos de alta movilidad.",
        "Sierra: reforzar atención a gravedad y respuesta oportuna.",
        "Viernes por la tarde: enfocar vigilancia, campañas y control operativo.",
        "Atropellamientos: proteger cruces, peatones y velocidad urbana.",
    ], 16)
    add_image(slide, paths["temporal_heat"], 6.8, 1.55, 5.55, 3.7)

    # 18
    slide = new_slide("Conclusiones", "La minería de datos permitió pasar de registros a evidencia accionable")
    add_bullets(slide, 0.9, 1.45, 10.9, 4.4, [
        "ATUS permitió analizar diez años de accidentes de Sonora con una base oficial.",
        "La preparación de datos fue clave para construir región, gravedad y variables temporales.",
        "El análisis confirmó concentración territorial y patrones por tipo, causa y horario.",
        "Los modelos apoyan la identificación de accidentes graves y la interpretación de variables relevantes.",
        "El valor del proyecto está en orientar prevención, no solo en describir datos.",
    ], 17)

    # 19
    slide = new_slide("Recomendaciones", "El trabajo futuro debe acercar el análisis a decisiones operativas")
    add_bullets(slide, 0.9, 1.45, 5.55, 4.1, [
        "Incorporar ubicación geográfica más precisa si está disponible.",
        "Cruzar con infraestructura vial, clima, tráfico y población.",
        "Actualizar el análisis cada año para monitorear tendencias.",
        "Diseñar tableros preventivos por región y municipio.",
        "Validar acciones con autoridades o especialistas en movilidad.",
    ], 16)
    add_textbox(slide, 7.0, 2.0, 4.65, 1.35, "Siguiente paso: convertir hallazgos en intervenciones medibles.", 26, ORANGE, True, PP_ALIGN.CENTER)

    # 20
    slide = new_slide("Referencias", "Fuentes utilizadas")
    refs = [
        "INEGI. Accidentes de Tránsito Terrestre en Zonas Urbanas y Suburbanas (ATUS), 2015-2024.",
        "E1_Propuesta/E1_Propuesta.pdf.",
        "E2_Analisis_Exploratorio_Preprocesamiento.ipynb.",
        "E3_Modelos_Evaluacion_ATUS_Sonora.ipynb.",
        "Bibliotecas: pandas, numpy, matplotlib, seaborn, scipy, scikit-learn y mlxtend.",
    ]
    add_bullets(slide, 0.95, 1.55, 11.2, 3.6, refs, 15)

    prs.save(PPTX_PATH)
    return slides_meta


def write_guide(analysis: dict):
    entries = [
        ("Portada", "Tema, fuente, periodo y enfoque preventivo.", "Gráfica de accidentes por año como apoyo visual.", "Presentar el proyecto y dejar claro que el objetivo no es solo describir, sino prevenir.", "Título grande, fondo oscuro y tres indicadores clave.", "Para entender por qué importa, pasamos al problema de estudio."),
        ("Problema de estudio", "Los accidentes viales generan daños, lesiones y fallecimientos.", "Gráfica de evolución anual.", "Explicar que el riesgo vial tiene patrones por lugar, tiempo y tipo de accidente.", "Dos columnas: problema a la izquierda y gráfica a la derecha.", "A partir del problema, planteamos una pregunta y un objetivo."),
        ("Pregunta de investigación y objetivo", "Pregunta guía y objetivo general del proyecto.", "Bloques de texto sin gráfica.", "Mostrar que la investigación busca detectar patrones de riesgo en datos ATUS.", "Dos bloques equilibrados: pregunta y objetivo.", "Después definimos la fuente de datos."),
        ("Fuente de datos ATUS / INEGI", "Datos oficiales ATUS, Sonora, 2015-2024.", "Indicadores de periodo, fuente y registros.", "Explicar que INEGI da una base confiable y comparable por año.", "Tarjetas superiores y una gráfica compacta.", "Con la fuente clara, explicamos cómo se limpió la base."),
        ("Preparación y limpieza de datos", "Unificación, filtro Sonora, eliminación de Certificado cero y variables nuevas.", "Tabla de pasos de limpieza.", "Subrayar que limpiar datos evita conclusiones equivocadas.", "Lista breve con tabla resumen.", "Luego se muestran las variables creadas."),
        ("Variables creadas: región, gravedad y víctimas", "REGION, TOTAL_VICTIMAS, GRAVE_BIN, RANGO_HORA e indicadores de moto/bici.", "Tabla de variables y uso.", "Explicar que estas variables traducen registros a riesgo vial.", "Tabla dominante y notas preventivas al lado.", "Con variables listas, pasamos al análisis exploratorio."),
        ("Análisis exploratorio general", "Crecimiento reciente y concentración territorial.", "Accidentes por año y por región.", "Mencionar el pico de 2023 y la concentración en regiones urbanas.", "Dos gráficas lado a lado con una frase final.", "Después hacemos zoom territorial."),
        ("Accidentes por región y municipio", "Hermosillo concentra volumen; Sierra destaca por gravedad.", "Gráfica regional y top de municipios.", "Aclarar que volumen y gravedad son dimensiones diferentes.", "Gráfica de región a la izquierda y municipios a la derecha.", "Luego vemos tipo de accidente y causa."),
        ("Accidentes por tipo y causa probable", "Predomina colisión con vehículo automotor y causa Conductor.", "Barras de tipo y causa.", "Conectar el hallazgo con velocidad, distancia, atención y señalamiento.", "Gráfica amplia con tres bullets de lectura.", "Después analizamos cuándo ocurren los riesgos."),
        ("Análisis temporal", "Años, meses y horarios críticos; tarde del viernes como ventana relevante.", "Heatmap día-hora y barras mensuales.", "Explicar que el tiempo ayuda a planear vigilancia y campañas específicas.", "Heatmap y barras por mes con llamada inferior.", "Luego medimos asociación estadística."),
        ("Chi-cuadrada y Cramér's V", "Medición de asociación entre variables y gravedad.", "Barras de Cramér's V y tabla de p-valores.", "Aclarar que asociación no significa causalidad, pero sí orienta prioridades.", "Gráfica principal y tabla breve.", "Después buscamos combinaciones frecuentes con Apriori."),
        ("Reglas de asociación con Apriori", "Combinaciones de condiciones que aparecen con accidentes graves.", "Tabla de reglas: soporte, confianza y lift.", "Explicar que Apriori ayuda a convertir patrones en mensajes preventivos.", "Tabla ancha con explicación debajo.", "Luego pasamos a modelos predictivos."),
        ("Modelos de clasificación", "Árbol, regresión logística y Random Forest para clasificar gravedad.", "Diagrama conceptual y tarjeta de Random Forest.", "Explicar variable objetivo GRAVE_BIN y evitar fuga de información.", "Lista técnica a la izquierda, idea clave a la derecha.", "Después comparamos el desempeño."),
        ("Comparación de modelos", "Métricas: accuracy, precisión, recall, F1 y ROC-AUC.", "Gráfica de métricas y tabla comparativa.", "Decir que en prevención importa detectar casos graves, no solo acertar en promedio.", "Gráfica grande y tabla pequeña.", "Luego interpretamos el modelo fuerte."),
        ("Importancia de variables con Random Forest", "Variables que más aportan a distinguir gravedad.", "Gráfica de importancia de variables.", "Explicar que la importancia ayuda a interpretar señales, no a probar causalidad.", "Gráfica horizontal y bullets de cautela.", "Después agrupamos municipios."),
        ("Clustering jerárquico de municipios", "Municipios agrupados por comportamiento similar.", "Dendrograma de clustering.", "Mostrar que permite comparar perfiles de riesgo y no solo municipios aislados.", "Dendrograma amplio con explicación lateral.", "Luego traducimos resultados a prevención."),
        ("Interpretación preventiva", "Acciones sugeridas por región, horario y tipo de riesgo.", "Heatmap temporal como evidencia.", "Enfatizar que los resultados deben convertirse en decisiones concretas.", "Lista de acciones y gráfica de apoyo.", "Después cerramos conclusiones."),
        ("Conclusiones", "De datos oficiales a evidencia accionable.", "Lista de conclusiones.", "Resumir E2 y E3: preparación, EDA, modelos y enfoque preventivo.", "Lista amplia y limpia.", "Luego proponemos trabajo futuro."),
        ("Recomendaciones", "Mejoras: georreferenciación, clima, tráfico, tableros y actualización anual.", "Bloque visual de siguiente paso.", "Explicar que el proyecto puede crecer hacia decisiones operativas.", "Bullets a la izquierda y frase fuerte a la derecha.", "Finalmente citamos fuentes."),
        ("Referencias", "INEGI, E1, E2, E3 y bibliotecas usadas.", "Lista de referencias.", "Cerrar mencionando que todo se basó en datos oficiales y notebooks reproducibles.", "Lista limpia sin saturar.", "Cierre de la exposición."),
    ]

    lines = ["# Guion de exposición\n\n"]
    for i, (title, slide_text, evidence, oral, visual, connection) in enumerate(entries, start=1):
        lines.extend([
            f"## Diapositiva {i}: {title}\n\n",
            f"1. **Título de la diapositiva:** {title}\n",
            f"2. **Texto breve en la diapositiva:** {slide_text}\n",
            f"3. **Gráfica, tabla o captura:** {evidence}\n",
            f"4. **Notas para exponer oralmente:** {oral}\n",
            f"5. **Recomendación visual:** {visual}\n",
            f"**Conexión con la siguiente:** {connection}\n\n",
        ])
    GUIDE_PATH.write_text("".join(lines), encoding="utf-8")


def main():
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)
    atus = load_data()
    paths = create_charts(atus)
    analysis = build_analysis(atus, paths)
    build_deck(atus, analysis, paths)
    write_guide(analysis)
    print(PPTX_PATH)
    print(GUIDE_PATH)


if __name__ == "__main__":
    main()
